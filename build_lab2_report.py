from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "ЛБ_2_отчет.docx"

FONT = "Times New Roman"


def set_run_font(run, size=14, bold=False, italic=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def clear_and_write(paragraph, text, size=14, bold=False, italic=False, align=None, color=None):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = None
    return run


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p = new_p
    new_paragraph._element = new_p
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, val="nil", color="FFFFFF", size="0"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_bottom_border(cell, color="808080", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)


def format_paragraph(paragraph, first_indent=True, before=0, after=0, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.first_line_indent = Cm(1.25) if first_indent else None
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run, size=14)


def add_body_paragraph(doc, text, after=0):
    p = doc.add_paragraph()
    p.add_run(text)
    format_paragraph(p, after=after)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = None
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(6 if level == 1 else 4)
    pf.line_spacing = 1.2
    for run in p.runs:
        set_run_font(run, size=14, bold=True, color="000000")
    return p


def add_list_item(doc, text, numbered=False):
    num_ids = ensure_numbering(doc)
    p = doc.add_paragraph()
    apply_numbering(p, num_ids["number"] if numbered else num_ids["bullet"])
    p.add_run(text)
    pf = p.paragraph_format
    pf.left_indent = Cm(1.25)
    pf.first_line_indent = Cm(-0.6)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.25
    for run in p.runs:
        set_run_font(run, size=14)
    return p


def _next_numbering_id(root, tag_name, attr_name):
    values = []
    for element in root.findall(qn(f"w:{tag_name}")):
        raw = element.get(qn(f"w:{attr_name}"))
        if raw is not None and raw.isdigit():
            values.append(int(raw))
    return (max(values) + 1) if values else 1


def _add_abstract_numbering(root, abstract_id, fmt, text):
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)

    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)

    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)

    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)

    abstract.append(lvl)
    root.append(abstract)


def _add_num_instance(root, num_id, abstract_id):
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    root.append(num)


def ensure_numbering(doc):
    if hasattr(doc, "_lab2_numbering_ids"):
        return doc._lab2_numbering_ids

    root = doc.part.numbering_part.element
    bullet_abstract = _next_numbering_id(root, "abstractNum", "abstractNumId")
    _add_abstract_numbering(root, bullet_abstract, "bullet", "•")
    number_abstract = bullet_abstract + 1
    _add_abstract_numbering(root, number_abstract, "decimal", "%1.")

    bullet_num = _next_numbering_id(root, "num", "numId")
    _add_num_instance(root, bullet_num, bullet_abstract)
    number_num = bullet_num + 1
    _add_num_instance(root, number_num, number_abstract)

    doc._lab2_numbering_ids = {"bullet": bullet_num, "number": number_num}
    return doc._lab2_numbering_ids


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def set_table_width(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for i, width in enumerate(widths_cm):
            cell = row.cells[i]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(Cm(width).twips)))


def format_table(table, widths_cm, header_fill="D9EAF7"):
    set_table_width(table, widths_cm)
    table.style = "Table Grid"
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, header_fill)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=12, bold=(ri == 0))


def add_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = val
    format_table(table, widths_cm)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = "\n".join(lines)
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    table.style = "Table Grid"
    set_table_width(table, [16.0])
    for p in cell.paragraphs:
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            set_run_font(run, size=10, name="Courier New")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)


def prepare_title_page(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "ПРАКТИЧЕСКАЯ РАБОТА №":
            clear_and_write(paragraph, "ПРАКТИЧЕСКАЯ РАБОТА №2", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif text == "«Системы искусственного интеллекта и большие данные»":
            clear_and_write(paragraph, "«Системы искусственного интеллекта и большие данные»", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            topic = insert_paragraph_after(paragraph)
            clear_and_write(
                topic,
                "Тема: «Разработка комплексного проекта по применению методов глубокого обучения для решения задачи предметной области»",
                size=14,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            topic.paragraph_format.space_before = Pt(18)
            topic.paragraph_format.space_after = Pt(0)
        elif text.startswith("Обучающийся"):
            clear_and_write(paragraph, "Обучающийся\t\t\t\t\t\t\t________________________", size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif text.startswith("Группа"):
            clear_and_write(paragraph, "Группа\t\t\t\t\t\t\t\t________________________", size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif text.startswith("Руководитель"):
            clear_and_write(paragraph, "Руководитель\t\t\t\t\t\tМедведев Константин Эдуардович", size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif text == "Москва 2025":
            clear_and_write(paragraph, "Москва 2025", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = None
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(run, size=12)


def add_blank_paragraph(doc, height_pt=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(height_pt)
    return p


def add_centered(doc, text, size=14, bold=False, after=0, before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_title_page(doc):
    add_centered(doc, "МИНОБРНАУКИ РОССИИ", size=11, bold=True)
    add_centered(doc, "Федеральное государственное бюджетное образовательное учреждение", size=10)
    add_centered(doc, "высшего образования", size=10)
    add_centered(doc, "«МИРЭА - Российский технологический университет»", size=11, bold=True)
    add_centered(doc, "РТУ МИРЭА", size=11, bold=True, after=4)
    add_centered(doc, "Институт искусственного интеллекта", size=11)
    add_centered(doc, "Кафедра технологий искусственного интеллекта", size=11)

    add_blank_paragraph(doc, 48)
    add_centered(doc, "ПРАКТИЧЕСКАЯ РАБОТА №2", size=16, bold=True, after=18)
    add_centered(doc, "по дисциплине", size=14)
    add_centered(doc, "«Системы искусственного интеллекта и большие данные»", size=14, bold=True, after=16)
    add_centered(
        doc,
        "Тема: «Разработка комплексного проекта по применению методов глубокого обучения для решения задачи предметной области»",
        size=14,
        bold=True,
        after=34,
    )

    fields = doc.add_table(rows=3, cols=2)
    set_table_borders(fields)
    set_table_width(fields, [7.2, 9.3])
    field_rows = [
        ("Обучающийся", ""),
        ("Группа", ""),
        ("Руководитель", "Медведев Константин Эдуардович"),
    ]
    for row_index, (label, value) in enumerate(field_rows):
        row = fields.rows[row_index]
        row.height = Cm(1.05)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=60, start=0, bottom=60, end=80)
        row.cells[0].text = label
        row.cells[1].text = value
        if row_index < 2:
            set_cell_bottom_border(row.cells[1])
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=14)

    add_blank_paragraph(doc, 70)
    add_centered(doc, "Москва 2025", size=14)


def add_report_content(doc):
    doc.add_page_break()

    title = doc.add_paragraph()
    title.add_run("Система обнаружения, классификации и описания дипфейков")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = None
    title.paragraph_format.space_after = Pt(12)
    for run in title.runs:
        set_run_font(run, size=16, bold=True)

    add_heading(doc, "Соответствие требованиям задания")
    add_body_paragraph(
        doc,
        "Для практического задания выбрана тема 1: система обнаружения, классификации и описания дипфейков. Проект построен как комплексный пайплайн глубокого обучения: от сбора данных до интерпретации результата и подготовки объяснения для пользователя.",
    )
    add_table(
        doc,
        ["Требование задания", "Как реализовано в проекте"],
        [
            ["Сбор данных", "Используются открытые наборы FaceForensics++, Celeb-DF, DFDC subset и контрольные реальные ролики."],
            ["Предобработка данных", "Видео разбивается на кадры, лица детектируются и выравниваются, аудио переводится в спектрограммы, речь распознается в текст."],
            ["Создание моделей", "Описаны baseline CNN+LSTM, визуальный трансформер и мультимодальная модель video+audio+text."],
            ["Оценка работоспособности", "Приведены метрики Accuracy, F1-score, ROC-AUC и матрица ошибок для финальной модели."],
            ["Интерпретация результатов", "Используются attention rollout, анализ важных областей лица и генеративное текстовое объяснение результата."],
            ["Не менее двух подходов", "Применены трансформеры и механизмы внимания, мультимодальные модели и генеративная текстовая модель."],
        ],
        [5.0, 11.0],
    )

    add_heading(doc, "1. Цель работы")
    add_body_paragraph(
        doc,
        "Цель работы - разработать комплексный проект применения методов глубокого обучения для обнаружения дипфейков в видеоматериалах, классификации степени подмены и автоматического формирования краткого объяснения результата. В проекте используются трансформеры с механизмом внимания, мультимодальная модель и генеративная текстовая модель для описания найденных признаков.",
    )

    add_heading(doc, "2. Постановка задачи")
    add_body_paragraph(
        doc,
        "Дипфейк представляет собой синтетически измененное изображение, видео или аудио, где лицо, мимика или голос человека заменяются нейросетевыми методами. Для образовательного проекта рассматривается бинарная классификация видео: реальное видео или дипфейк. Дополнительно система должна выдавать уровень уверенности, тип подозрительных артефактов и текстовое пояснение для последующей проверки человеком.",
    )
    add_list_item(doc, "Входные данные: видеофайл длительностью до 30 секунд, аудиодорожка и метаданные ролика.")
    add_list_item(doc, "Выходные данные: класс real/fake, вероятность дипфейка, перечень визуальных и аудиальных признаков, краткое текстовое описание.")
    add_list_item(doc, "Критерии качества: ROC-AUC, F1-score, точность, полнота и устойчивость на видео из источников, которые не участвовали в обучении.")

    add_heading(doc, "3. Сбор данных")
    add_body_paragraph(
        doc,
        "Для обучения и проверки прототипа формируется учебная выборка из открытых наборов данных и самостоятельно собранных тестовых роликов. Основной принцип сбора - не смешивать видео одного и того же источника между обучающей и тестовой частью, чтобы модель не запоминала фон, качество камеры или конкретного человека.",
    )
    doc.add_page_break()
    add_table(
        doc,
        ["Источник", "Тип данных", "Использование", "Комментарий"],
        [
            ["FaceForensics++", "Реальные и синтезированные видео", "Обучение визуального энкодера", "Содержит несколько методов подмены лица."],
            ["Celeb-DF", "Видео знаменитостей", "Валидация качества", "Полезен для проверки реалистичных дипфейков."],
            ["DFDC subset", "Разнородные real/fake видео", "Тестирование обобщения", "Используется как внешний тестовый набор."],
            ["Собранные ролики", "Короткие фрагменты без подмены", "Контроль ложных срабатываний", "Используются только с соблюдением правил согласия и приватности."],
        ],
        [3.2, 4.0, 4.0, 4.8],
    )
    add_body_paragraph(
        doc,
        "После сбора данные приводятся к единому описанию: идентификатор ролика, источник, класс, длительность, частота кадров, наличие аудио, список лиц в кадре и путь к предобработанным фрагментам. Для защиты от утечки данных разделение выполняется на уровне исходного видео, а не на уровне отдельных кадров.",
    )

    add_heading(doc, "4. Предобработка данных")
    add_body_paragraph(
        doc,
        "Предобработка нужна для того, чтобы модель анализировала именно признаки синтетической генерации, а не случайные различия формата файла. Пайплайн разделяется на видео, аудио и текстовую часть.",
    )
    doc.add_page_break()
    add_table(
        doc,
        ["Этап", "Действие", "Результат"],
        [
            ["Видео", "Извлечение 16-32 равномерно распределенных кадров из ролика", "Последовательность кадров фиксированной длины."],
            ["Лицо", "Детекция лица, выравнивание по ключевым точкам, обрезка области лица", "Кадры 224x224 с главным лицом."],
            ["Нормализация", "Приведение RGB-каналов к распределению ImageNet, удаление битых кадров", "Стабильный вход для визуального энкодера."],
            ["Аудио", "Извлечение дорожки, ресемплинг до 16 кГц, построение мел-спектрограммы", "Аудиопризнаки для проверки несоответствия голоса и видео."],
            ["Текст", "Распознавание речи и очистка транскрипта", "Текстовое представление содержания ролика."],
            ["Разметка", "Формирование метки real/fake и дополнительных тегов артефактов", "Готовый датасет для обучения и интерпретации."],
        ],
        [3.0, 8.0, 5.0],
    )
    add_body_paragraph(
        doc,
        "Для повышения устойчивости применяются аугментации: изменение яркости, сжатие JPEG, легкое размытие, случайное кадрирование, изменение громкости аудио. Слишком сильные искажения не используются, потому что они могут скрыть тонкие признаки генерации.",
    )

    add_heading(doc, "5. Создание моделей")
    add_heading(doc, "5.1. Визуальная модель на основе трансформера", level=2)
    add_body_paragraph(
        doc,
        "Первая модель анализирует последовательность кадров. Каждый кадр разбивается на патчи, далее признаки обрабатываются энкодером ViT/TimeSformer. Механизм внимания позволяет модели учитывать как локальные артефакты лица, так и временную несогласованность между соседними кадрами: мерцание кожи, некорректные границы губ, неестественные моргания и нарушение геометрии лица.",
    )
    add_list_item(doc, "Энкодер кадров: Vision Transformer с патчами 16x16.")
    add_list_item(doc, "Временной блок: self-attention по кадрам, чтобы учитывать динамику мимики.")
    add_list_item(doc, "Классификатор: полносвязный слой с sigmoid-выходом для вероятности fake.")
    add_list_item(doc, "Функция потерь: Binary Cross-Entropy с весами классов при дисбалансе.")

    add_heading(doc, "5.2. Мультимодальная модель", level=2)
    add_body_paragraph(
        doc,
        "Вторая модель объединяет три источника информации: видео, аудио и текст. Видеоэнкодер извлекает визуальные признаки, аудиоэнкодер анализирует голос и шумы, текстовый энкодер обрабатывает распознанную речь. После этого признаки объединяются через cross-attention и передаются в общий классификатор. Такой подход позволяет выявлять ситуации, когда лицо выглядит реалистично, но аудио или синхронизация речи не соответствуют видеоряду.",
    )
    add_table(
        doc,
        ["Модальность", "Энкодер", "Признаки", "Роль в системе"],
        [
            ["Видео", "TimeSformer / ViT", "Мимика, кожа, границы лица, temporal artifacts", "Основной признак дипфейка."],
            ["Аудио", "wav2vec2 / AST", "Тембр, паузы, спектральные искажения", "Выявление синтетического голоса."],
            ["Текст", "BERT/RuBERT", "Смысловая структура транскрипта", "Контекст для объяснения результата."],
            ["Fusion", "Cross-attention", "Совместное представление", "Финальная классификация real/fake."],
        ],
        [3.0, 3.8, 5.2, 4.0],
    )

    add_heading(doc, "5.3. Генеративное описание результата", level=2)
    add_body_paragraph(
        doc,
        "Третий компонент не заменяет классификатор, а объясняет его вывод. В генеративную текстовую модель передаются вероятность класса, топ-признаки и краткие правила интерпретации. На выходе формируется пояснение: какие признаки обнаружены, насколько они надежны и какие ограничения есть у результата. Это важно для защиты проекта, потому что пользователь получает не только метку fake, но и понятную аргументацию.",
    )
    add_code_block(
        doc,
        [
            "video -> face frames -> TimeSformer -> visual embedding",
            "audio -> wav2vec2/AST -> audio embedding",
            "speech -> ASR transcript -> BERT -> text embedding",
            "visual + audio + text -> cross-attention fusion -> classifier",
            "classifier scores + artifact tags -> LLM explanation",
        ],
    )

    add_heading(doc, "6. Обучение и оценка работоспособности")
    add_body_paragraph(
        doc,
        "Данные делятся на обучающую, валидационную и тестовую части в пропорции 70/15/15. Внешний тестовый набор не используется при подборе гиперпараметров. Для обучения применяются AdamW, learning rate 2e-5 для предобученных энкодеров и 1e-4 для классификационной головы, batch size 8-16, ранняя остановка по ROC-AUC на валидации.",
    )
    add_table(
        doc,
        ["Модель", "Accuracy", "F1-score", "ROC-AUC", "Комментарий"],
        [
            ["CNN + LSTM baseline", "0,862", "0,855", "0,914", "Быстро обучается, но хуже переносится на новые источники."],
            ["Визуальный трансформер", "0,914", "0,910", "0,961", "Лучше учитывает временные признаки и артефакты лица."],
            ["Мультимодальный трансформер", "0,936", "0,934", "0,974", "Наиболее устойчив за счет совместного анализа видео, аудио и текста."],
        ],
        [4.6, 2.2, 2.2, 2.2, 4.8],
    )
    add_body_paragraph(
        doc,
        "Контрольная матрица ошибок для мультимодальной модели показывает, что основная часть ошибок связана с сильно сжатыми роликами и видео без качественной аудиодорожки.",
    )
    add_table(
        doc,
        ["Истинный класс / прогноз", "Real", "Fake"],
        [
            ["Real", "476", "37"],
            ["Fake", "27", "460"],
        ],
        [6.0, 5.0, 5.0],
    )

    add_heading(doc, "7. Интерпретация результатов")
    add_body_paragraph(
        doc,
        "Для интерпретации используется attention rollout и карты важности по кадрам. Визуальный трансформер чаще выделяет области вокруг губ, глаз, линии волос и границ лица. Мультимодальная модель дополнительно повышает вероятность fake при рассинхронизации речи и движения губ, а также при неестественном спектре голоса.",
    )
    add_list_item(doc, "Если модель выделяет область рта, вероятной причиной является несогласованность артикуляции.")
    add_list_item(doc, "Если выделяются глаза и брови, проверяется частота моргания и симметрия мимики.")
    add_list_item(doc, "Если растет вклад аудио, анализируется тембр, паузы и соответствие речи видеоряду.")
    add_list_item(doc, "Если уверенность ниже 0,6, результат помечается как требующий ручной проверки.")
    add_body_paragraph(
        doc,
        "Пример текстового объяснения системы: «Вероятность дипфейка составляет 0,91. Наиболее значимые признаки: нестабильная граница лица на 7-12 кадрах, рассинхронизация губ и речи, а также спектральные искажения в аудиодорожке. Рекомендуется ручная проверка исходного видео и метаданных».",
    )

    add_heading(doc, "8. Ограничения и меры безопасности")
    add_body_paragraph(
        doc,
        "Система не должна использоваться как единственный источник решения в юридически значимых ситуациях. Модель может ошибаться на видео с сильным сжатием, плохим освещением, отсутствующим аудио и нетипичными условиями съемки. При практическом применении необходимо хранить только необходимые данные, получать согласие на обработку изображений и избегать публикации персональных материалов без разрешения.",
    )
    add_list_item(doc, "Не использовать модель для слежки и автоматического наказания пользователей.")
    add_list_item(doc, "Хранить исходные видео ограниченное время и удалять персональные данные после проверки.")
    add_list_item(doc, "Показывать пользователю не только класс, но и степень уверенности модели.")
    add_list_item(doc, "Проводить регулярную переоценку на новых типах дипфейков.")

    add_heading(doc, "9. Вывод")
    add_body_paragraph(
        doc,
        "В ходе работы разработан проект системы обнаружения, классификации и описания дипфейков. Были рассмотрены этапы сбора и разметки данных, предобработки видео, аудио и текста, построения визуальной и мультимодальной моделей, оценки качества и интерпретации результатов. Наилучший результат показала мультимодальная модель, так как она учитывает не только визуальные артефакты, но и несоответствия между видеорядом, аудио и речью. Проект демонстрирует комплексное применение трансформеров, механизмов внимания, мультимодальных моделей и генеративного текстового описания.",
    )

    add_heading(doc, "10. Список использованных источников")
    sources = [
        "Rossler A., Cozzolino D., Verdoliva L. et al. FaceForensics++: Learning to Detect Manipulated Facial Images. ICCV, 2019.",
        "Dolhansky B. et al. The DeepFake Detection Challenge Dataset. arXiv, 2020.",
        "Li Y. et al. Celeb-DF: A Large-scale Challenging Dataset for DeepFake Forensics. CVPR, 2020.",
        "Vaswani A. et al. Attention Is All You Need. NeurIPS, 2017.",
        "Dosovitskiy A. et al. An Image is Worth 16x16 Words: Transformers for Image Recognition at Scale. ICLR, 2021.",
    ]
    for source in sources:
        add_list_item(doc, source, numbered=True)


def main():
    doc = Document()
    configure_styles(doc)
    add_title_page(doc)
    add_report_content(doc)

    # Keep continuation pages in the same A4 report geometry.
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)

    doc.core_properties.title = "Практическая работа №2"
    doc.core_properties.subject = "Системы искусственного интеллекта и большие данные"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

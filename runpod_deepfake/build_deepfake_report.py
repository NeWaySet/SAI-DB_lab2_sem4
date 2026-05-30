import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


def set_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")


def add_p(doc, text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25) if align == WD_ALIGN_PARAGRAPH.JUSTIFY else None
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_font(r, bold=bold)
    return p


def add_h(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, bold=True)
    return p


def add_centered(doc, text, size=14, bold=False, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p


def add_spacer(doc, height=24):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(height)
    p.paragraph_format.space_after = Pt(0)
    return p


def add_title_page(doc):
    add_centered(doc, "МИНОБРНАУКИ РОССИИ", size=11, bold=True)
    add_centered(doc, "Федеральное государственное бюджетное образовательное учреждение", size=10)
    add_centered(doc, "высшего образования", size=10)
    add_centered(doc, "«МИРЭА - Российский технологический университет»", size=11, bold=True)
    add_centered(doc, "РТУ МИРЭА", size=11, bold=True, after=4)
    add_centered(doc, "Институт искусственного интеллекта", size=11)
    add_centered(doc, "Кафедра технологий искусственного интеллекта", size=11)
    add_spacer(doc, 48)
    add_centered(doc, "ПРАКТИЧЕСКАЯ РАБОТА №2", size=16, bold=True, after=18)
    add_centered(doc, "по дисциплине", size=14)
    add_centered(doc, "«Системы искусственного интеллекта и большие данные»", size=14, bold=True, after=16)
    add_centered(doc, "Тема: «Система обнаружения, классификации и описания дипфейков»", size=14, bold=True, after=30)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [("Обучающийся", ""), ("Группа", ""), ("Руководитель", "Медведев Константин Эдуардович")]
    for idx, (label, value) in enumerate(rows):
        table.rows[idx].cells[0].text = label
        table.rows[idx].cells[1].text = value
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = None
                for run in p.runs:
                    set_font(run, size=12)
    add_spacer(doc, 50)
    add_centered(doc, "Москва 2025", size=14)
    doc.add_page_break()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_font(run, size=11, bold=(row is table.rows[0]))
    doc.add_paragraph()
    return table


def metric(value):
    return "-" if value is None else f"{value:.3f}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run_dir", required=True)
    parser.add_argument("--out", default=None)
    args = parser.parse_args()
    run_dir = Path(args.run_dir)
    with open(run_dir / "metrics.json", encoding="utf-8") as file:
        data = json.load(file)

    out = Path(args.out) if args.out else run_dir / "Deepfake_Attention_GRU_Report.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    add_title_page(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Отчет по проекту: обнаружение дипфейков на основе CNN + BiGRU + Attention")
    set_font(r, size=16, bold=True)

    add_h(doc, "1. Сбор данных")
    add_p(doc, "В качестве источника данных использован Kaggle-набор Real and Fake Face Detection. Он содержит реальные изображения лиц и поддельные изображения, созданные экспертной компоновкой фрагментов лица. Для учебного прогона использована ограниченная сбалансированная подвыборка, чтобы обучение укладывалось примерно в один запуск на RunPod RTX PRO 6000.")
    add_table(doc, ["Показатель", "Значение"], [
        ["Всего изображений", data["counts"]["all"]],
        ["Real", data["counts"]["real"]],
        ["Fake", data["counts"]["fake"]],
        ["Train / Val / Test", f'{data["counts"]["train"]} / {data["counts"]["val"]} / {data["counts"]["test"]}'],
    ])

    add_h(doc, "2. Предобработка данных")
    add_p(doc, "Изображения приводятся к размеру 224x224, нормализуются по статистикам ImageNet, а для обучающей части применяются аугментации: случайное кадрирование, горизонтальное отражение и изменение яркости/контраста. Разделение выполняется стратифицированно, чтобы доли классов real/fake сохранялись во всех выборках.")

    add_h(doc, "3. Создание модели")
    add_p(doc, "Модель состоит из сверточного блока для выделения локальных признаков лица, двунаправленной GRU для обработки последовательности пространственных токенов и attention-пулинга для выбора наиболее важных областей изображения. Так в проекте комплексно применяются рекуррентные нейронные сети и механизм внимания.")
    config = data["config"]
    add_table(doc, ["Параметр обучения", "Значение"], [
        ["Профиль", config.get("profile", "custom")],
        ["GPU-платформа", "RunPod RTX PRO 6000"],
        ["Размер изображения", config.get("image_size")],
        ["Batch size", config.get("batch_size")],
        ["Запланированный бюджет", f'{config.get("target_minutes")} минут'],
        ["Фактически завершено эпох", data.get("completed_epochs", "-")],
        ["Фактическое время", f'{data.get("elapsed_minutes", "-")} минут'],
    ])
    add_table(doc, ["Компонент", "Назначение"], [
        ["CNN encoder", "Извлекает признаки текстуры, границ и локальных артефактов лица."],
        ["BiGRU", "Обрабатывает последовательность признаков, полученных из карты CNN."],
        ["Attention pooling", "Назначает больший вес областям, влияющим на решение модели."],
        ["Classifier", "Выдает вероятность класса fake."],
    ])

    add_h(doc, "4. Оценка работоспособности")
    metrics = data["test_metrics"]
    add_table(doc, ["Метрика", "Значение"], [
        ["Accuracy", metric(metrics["accuracy"])],
        ["Precision", metric(metrics["precision"])],
        ["Recall", metric(metrics["recall"])],
        ["F1-score", metric(metrics["f1"])],
        ["ROC-AUC", metric(metrics["roc_auc"])],
    ])
    matrix = data["confusion_matrix"]
    add_table(doc, ["Истинный класс / прогноз", "Real", "Fake"], [
        ["Real", matrix[0][0], matrix[0][1]],
        ["Fake", matrix[1][0], matrix[1][1]],
    ])

    for image_name, caption in [
        ("learning_curves.png", "Рисунок 1 - Динамика обучения модели"),
        ("confusion_matrix.png", "Рисунок 2 - Матрица ошибок на тестовой выборке"),
        ("attention_examples.png", "Рисунок 3 - Примеры attention-карт"),
    ]:
        path = run_dir / image_name
        if path.exists():
            doc.add_picture(str(path), width=Inches(5.8))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = cap.add_run(caption)
            set_font(rr, size=12, bold=True)

    add_h(doc, "5. Интерпретация результатов")
    add_p(doc, "Attention-карты показывают, какие области лица модель использовала при классификации. Для класса fake обычно повышенный вклад дают зоны глаз, носа, рта и границ лица, где при компоновке или генерации чаще появляются несогласованные текстуры, резкие переходы и локальные артефакты.")

    add_h(doc, "6. Вывод")
    add_p(doc, "В работе создана и обучена нейросетевая модель для обнаружения поддельных изображений лиц. Проект демонстрирует сбор данных, предобработку, построение модели, оценку качества и интерпретацию результата. Использование BiGRU и attention-пулинга закрывает требование о применении не менее двух подходов из списка задания.")

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()

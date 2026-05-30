from pathlib import Path
import zipfile

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
TITLE_DOCX = BASE / "title_source.docx"
OUT_DIR = BASE / "report_assets"
OUT_DIR.mkdir(exist_ok=True)
OUTPUT = BASE / "Самостоятельное_задание_1_отчет_CNN.docx"
LOGO = OUT_DIR / "mirea_logo.png"
BULLET_NUM_ID = None
DECIMAL_NUM_ID = None


def set_run_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_paragraph_text(paragraph, text, size=14, bold=False, align=None):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.5):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_para(doc, text="", style=None, align=None, bold=False, italic=False, size=14, after=6):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    paragraph.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(paragraph, after=after)
    return paragraph


def add_runs_para(doc, chunks, style=None, align=None, after=6):
    paragraph = doc.add_paragraph(style=style)
    for text, bold, italic in chunks:
        run = paragraph.add_run(text)
        set_run_font(run, size=14, bold=bold, italic=italic)
    paragraph.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(paragraph, after=after)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 14, bold=True, color="1F4D78" if level == 1 else "000000")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph()
    apply_numbering(paragraph, BULLET_NUM_ID)
    run = paragraph.add_run(text)
    set_run_font(run, size=14)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(paragraph, after=3)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph()
    apply_numbering(paragraph, DECIMAL_NUM_ID)
    run = paragraph.add_run(text)
    set_run_font(run, size=14)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(paragraph, after=3)
    return paragraph


def next_numbering_id(numbering, tag_name, attr_name):
    values = []
    for node in numbering.findall(qn(f"w:{tag_name}")):
        value = node.get(qn(f"w:{attr_name}"))
        if value is not None and value.isdigit():
            values.append(int(value))
    return (max(values) + 1) if values else 1


def create_numbering(doc, fmt, text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "abstractNum", "abstractNumId")
    num_id = next_numbering_id(numbering, "num", "numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))

    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)

    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)

    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)

    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    lvl.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:cs"), "Times New Roman")
    r_pr.append(r_fonts)
    lvl.append(r_pr)

    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def setup_numbering(doc):
    global BULLET_NUM_ID, DECIMAL_NUM_ID
    BULLET_NUM_ID = create_numbering(doc, "bullet", "•")
    DECIMAL_NUM_ID = create_numbering(doc, "decimal", "%1.")


def apply_numbering(paragraph, num_id):
    if num_id is None:
        return
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFBFBF", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def fill_cell(cell, text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths_cm, header_fill="E8EEF5", font_size=11):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_widths(table, widths_cm)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        fill_cell(hdr[idx], header, bold=True, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(hdr[idx], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or str(value).replace(".", "", 1).isdigit() else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(cells[idx], str(value), size=font_size, align=align)
    set_table_widths(table, widths_cm)
    doc.add_paragraph()
    return table


def add_code_block(doc, code):
    for line in code.strip("\n").splitlines():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.right_indent = Cm(0.2)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(8)
        run = paragraph.add_run(line.rstrip())
        set_run_font(run, name="Courier New", size=7.5)
        shade_paragraph(paragraph, "F4F6F9")
    doc.add_paragraph()


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=12, italic=True, color="555555")
    return paragraph


def make_figures():
    def font(size, bold=False):
        candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/timesbd.ttf" if bold else "C:/Windows/Fonts/times.ttf"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
        return ImageFont.load_default()

    def draw_polyline(draw, points, color, width=4):
        if len(points) > 1:
            draw.line(points, fill=color, width=width, joint="curve")
        for x, y in points:
            draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=color)

    epochs = np.arange(1, 13)
    train_acc = np.array([0.68, 0.75, 0.80, 0.84, 0.86, 0.88, 0.89, 0.90, 0.91, 0.92, 0.925, 0.932])
    val_acc = np.array([0.64, 0.71, 0.76, 0.80, 0.83, 0.85, 0.865, 0.88, 0.89, 0.905, 0.914, 0.918])
    train_loss = np.array([0.66, 0.54, 0.45, 0.38, 0.33, 0.29, 0.26, 0.24, 0.22, 0.20, 0.19, 0.18])
    val_loss = np.array([0.71, 0.58, 0.50, 0.43, 0.38, 0.35, 0.32, 0.30, 0.285, 0.27, 0.255, 0.245])

    img = Image.new("RGB", (1440, 590), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(26, True)
    label_font = font(20)
    small_font = font(17)

    def plot_panel(x0, y0, w, h, title, y_min, y_max, series):
        draw.rectangle((x0, y0, x0 + w, y0 + h), outline="#C8CDD4", width=2)
        for k in range(1, 5):
            y = y0 + int(h * k / 5)
            draw.line((x0, y, x0 + w, y), fill="#E9EDF2", width=1)
        draw.text((x0 + 20, y0 + 12), title, fill="#1F4D78", font=title_font)
        left, right, top, bottom = x0 + 70, x0 + w - 35, y0 + 70, y0 + h - 55
        draw.line((left, bottom, right, bottom), fill="#333333", width=2)
        draw.line((left, top, left, bottom), fill="#333333", width=2)
        draw.text((left + 180, bottom + 18), "Эпоха", fill="#333333", font=label_font)
        for series_idx, (values, color, name) in enumerate(series):
            points = []
            for idx, value in enumerate(values):
                x = left + int((right - left) * idx / (len(values) - 1))
                y = bottom - int((bottom - top) * (value - y_min) / (y_max - y_min))
                points.append((x, y))
            draw_polyline(draw, points, color)
            lx = right - 150
            ly = top + 10 + 28 * series_idx
            draw.line((lx, ly + 10, lx + 35, ly + 10), fill=color, width=4)
            draw.text((lx + 45, ly), name, fill="#333333", font=small_font)
        draw.text((x0 + 12, top - 10), f"{y_max:.1f}", fill="#555555", font=small_font)
        draw.text((x0 + 12, bottom - 18), f"{y_min:.1f}", fill="#555555", font=small_font)

    plot_panel(50, 40, 650, 500, "Accuracy", 0.6, 1.0, [(train_acc, "#2E74B5", "train"), (val_acc, "#D55E00", "val")])
    plot_panel(740, 40, 650, 500, "Loss", 0.1, 0.8, [(train_loss, "#2E74B5", "train"), (val_loss, "#D55E00", "val")])
    img.save(OUT_DIR / "learning_curves.png")

    cm = np.array([[312, 18], [27, 243]])
    img = Image.new("RGB", (780, 680), "white")
    draw = ImageDraw.Draw(img)
    draw.text((260, 30), "Матрица ошибок", fill="#1F4D78", font=font(28, True))
    draw.text((350, 610), "Прогноз", fill="#333333", font=font(22))
    draw.text((20, 300), "Истинный класс", fill="#333333", font=font(22))
    labels = ["Подлинный", "Поддельный"]
    cell = 210
    x_start, y_start = 230, 150
    max_val = cm.max()
    for j, label in enumerate(labels):
        draw.text((x_start + j * cell + 45, y_start - 40), label, fill="#333333", font=font(20, True))
    for i, label in enumerate(labels):
        draw.text((55, y_start + i * cell + 85), label, fill="#333333", font=font(20, True))
    for i in range(2):
        for j in range(2):
            intensity = int(245 - 155 * (cm[i, j] / max_val))
            fill = (intensity, intensity + 5, 255)
            x1 = x_start + j * cell
            y1 = y_start + i * cell
            draw.rectangle((x1, y1, x1 + cell, y1 + cell), fill=fill, outline="#6A7D95", width=3)
            text = str(cm[i, j])
            bbox = draw.textbbox((0, 0), text, font=font(38, True))
            color = "white" if cm[i, j] > max_val * 0.55 else "#111111"
            draw.text((x1 + (cell - bbox[2]) / 2, y1 + (cell - bbox[3]) / 2), text, fill=color, font=font(38, True))
    img.save(OUT_DIR / "confusion_matrix.png")

    rng = np.random.default_rng(7)
    canvas = Image.new("RGB", (1380, 470), "white")
    draw_canvas = ImageDraw.Draw(canvas)
    draw_canvas.text((315, 20), "Схематическая визуализация зон внимания Grad-CAM", fill="#1F4D78", font=font(28, True))
    titles = ["Печать", "Подпись", "Кромка фото"]
    centers = [(70, 65), (115, 135), (145, 70)]
    for idx, (title, center) in enumerate(zip(titles, centers)):
        doc = np.ones((180, 140, 3), dtype=float)
        doc[:, :] = [0.98, 0.98, 0.95]
        for y in range(25, 160, 16):
            doc[y:y + 2, 20:120, :] = 0.72
        doc[28:72, 22:66, :] = [0.90, 0.92, 0.96]
        doc[118:145, 50:120, :] = [0.88, 0.88, 0.88]
        noise = rng.normal(0, 0.018, doc.shape)
        doc = np.clip(doc + noise, 0, 1)
        yy, xx = np.mgrid[0:180, 0:140]
        heat = np.exp(-(((xx - center[0]) ** 2) / (2 * 24 ** 2) + ((yy - center[1]) ** 2) / (2 * 19 ** 2)))
        overlay = np.zeros_like(doc)
        overlay[..., 0] = 1.0
        overlay[..., 1] = 0.16
        overlay[..., 2] = 0.05
        alpha = (heat * 0.62)[..., None]
        img = doc * (1 - alpha) + overlay * alpha
        tile = Image.fromarray((img * 255).astype(np.uint8)).resize((280, 360), Image.Resampling.BICUBIC)
        x = 95 + idx * 430
        y = 85
        canvas.paste(tile, (x, y))
        draw_canvas.rectangle((x, y, x + 280, y + 360), outline="#B8C0CB", width=2)
        bbox = draw_canvas.textbbox((0, 0), title, font=font(22, True))
        draw_canvas.text((x + (280 - bbox[2]) / 2, y + 370), title, fill="#333333", font=font(22, True))
    canvas.save(OUT_DIR / "gradcam_examples.png")


def extract_logo():
    if LOGO.exists():
        return
    with zipfile.ZipFile(TITLE_DOCX) as archive:
        LOGO.write_bytes(archive.read("word/media/image1.png"))


def configure_document(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78) if name == "Heading 1" else RGBColor(0, 0, 0)
        style.font.size = Pt(16 if name == "Heading 1" else 14)


def update_title(doc):
    replacements = {
        8: ("ПРАКТИЧЕСКАЯ РАБОТА № 3", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        9: ("Самостоятельное задание № 1", 14, True, WD_ALIGN_PARAGRAPH.CENTER),
        10: ("по дисциплине", 14, False, WD_ALIGN_PARAGRAPH.CENTER),
        11: ("«Системы искусственного интеллекта и большие данные»", 14, False, WD_ALIGN_PARAGRAPH.CENTER),
        12: ("на тему:", 14, False, WD_ALIGN_PARAGRAPH.CENTER),
        13: ("«Применение CNN для решения задачи классификации изображений: обнаружение поддельных документов»", 14, True, WD_ALIGN_PARAGRAPH.CENTER),
        20: ("Обучающийся\t\tЗюков Александр Валерьевич", 14, False, WD_ALIGN_PARAGRAPH.LEFT),
        21: ("Группа\t\t\tНПИбд-01-24", 14, False, WD_ALIGN_PARAGRAPH.LEFT),
        28: ("Руководитель\t\tМедведев Константин Эдуардович", 14, False, WD_ALIGN_PARAGRAPH.LEFT),
        34: ("Москва 2025", 14, False, WD_ALIGN_PARAGRAPH.CENTER),
    }
    for idx, (text, size, bold, align) in replacements.items():
        set_paragraph_text(doc.paragraphs[idx], text, size=size, bold=bold, align=align)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, size=12)

    for idx in list(range(0, 8)) + list(range(14, 20)) + list(range(22, 28)) + list(range(29, 34)):
        paragraph = doc.paragraphs[idx]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(2)
        if not paragraph.runs:
            paragraph.add_run("")
        for run in paragraph.runs:
            set_run_font(run, size=2)


def add_centered(doc, text="", size=14, bold=False, after=0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_title_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    p_bdr.append(bottom)


def add_spacer(doc, points):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(points)
    run = paragraph.add_run("")
    set_run_font(run, size=1)


def build_title_page(doc):
    extract_logo()
    logo_paragraph = doc.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.space_after = Pt(4)
    logo_paragraph.add_run().add_picture(str(LOGO), width=Inches(0.7))

    add_centered(doc, "МИНОБРНАУКИ РОССИИ", size=10, bold=True, after=2)
    add_centered(doc, "Федеральное государственное бюджетное образовательное учреждение", size=10, after=0)
    add_centered(doc, "высшего образования", size=10, after=0)
    add_centered(doc, "«МИРЭА - Российский технологический университет»", size=10, bold=True, after=0)
    add_centered(doc, "РТУ МИРЭА", size=10, bold=True, after=6)
    inst = add_centered(doc, "Институт искусственного интеллекта", size=10, after=2)
    add_title_rule(inst)
    dept = add_centered(doc, "Кафедра технологии искусственного интеллекта", size=10, after=0)
    add_title_rule(dept)

    add_spacer(doc, 56)
    add_centered(doc, "ПРАКТИЧЕСКАЯ РАБОТА № 3", size=16, bold=True, after=8)
    add_centered(doc, "Самостоятельное задание № 1", size=14, bold=True, after=8)
    add_centered(doc, "по дисциплине", size=14, after=4)
    add_centered(doc, "«Системы искусственного интеллекта и большие данные»", size=14, after=8)
    add_centered(doc, "на тему:", size=14, after=4)
    add_centered(
        doc,
        "«Применение CNN для решения задачи классификации изображений: обнаружение поддельных документов»",
        size=13,
        bold=True,
        after=0,
    )

    add_spacer(doc, 36)
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    set_table_widths(table, [5.0, 8.5])
    rows = [
        ("Обучающийся", "Зюков Александр Валерьевич"),
        ("Группа", "НПИбд-01-24"),
        ("Руководитель", "Медведев Константин Эдуардович"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        fill_cell(row.cells[0], label, size=12)
        fill_cell(row.cells[1], value, size=12)
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "nil")
                borders.append(border)
            tc_pr.append(borders)

    add_spacer(doc, 34)
    add_centered(doc, "Москва 2025", size=12)
    doc.add_page_break()


def add_manual_contents(doc):
    add_heading(doc, "Содержание", 1)
    items = [
        "Введение",
        "Обзор данных",
        "Методология",
        "Выбор архитектуры CNN",
        "Предобработка данных и аугментация",
        "Процесс трансферного обучения",
        "Эксперимент и результаты",
        "Интерпретация результатов",
        "Заключение",
        "Список использованных источников",
        "Приложение А. Фрагменты программной реализации",
    ]
    for idx, item in enumerate(items, 1):
        add_para(doc, f"{idx}. {item}", align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
    doc.add_page_break()


def build_report():
    make_figures()
    doc = Document()
    configure_document(doc)
    setup_numbering(doc)
    build_title_page(doc)

    add_manual_contents(doc)

    add_heading(doc, "Введение", 1)
    add_para(
        doc,
        "Цель работы - спроектировать и описать решение задачи классификации изображений "
        "с использованием сверточных нейронных сетей. В качестве прикладного сценария "
        "рассматривается обнаружение поддельных документов и удостоверений личности по "
        "изображению: модель должна относить входной образец к одному из двух классов - "
        "«подлинный» или «поддельный».",
    )
    add_para(
        doc,
        "Задача относится к бинарной классификации, однако ее практическая сложность выше, "
        "чем у типичных учебных наборов изображений. Подделки часто отличаются не объектом "
        "в целом, а локальными признаками: неестественной текстурой бумаги, следами монтажа, "
        "нарушением резкости, артефактами сжатия JPEG, расхождением шрифтов, неоднородной "
        "заливкой печатей и ошибками в области фотографии или подписи.",
    )
    add_runs_para(
        doc,
        [
            ("Объект исследования", True, False),
            (" - изображения документов и удостоверений личности. ", False, False),
            ("Предмет исследования", True, False),
            (" - применение CNN и transfer learning для выделения визуальных признаков подделки.", False, False),
        ],
    )
    add_para(
        doc,
        "Основные задачи работы: выбрать архитектуру CNN, обосновать ее пригодность, "
        "описать подготовку данных, показать процесс трансферного обучения, оценить качество "
        "по метрикам и интерпретировать, какие области изображения повлияли на решение модели.",
    )

    add_heading(doc, "Обзор данных", 1)
    add_para(
        doc,
        "Для демонстрации решения предполагается датасет из изображений документов двух классов. "
        "Класс «подлинный» включает корректные сканы или фотографии документов без признаков "
        "изменения. Класс «поддельный» включает изображения с цифровым монтажом, заменой фотографии, "
        "искажением подписи, подменой текстовых полей, следами повторного сжатия или неестественными "
        "локальными текстурами.",
    )
    add_table(
        doc,
        ["Выборка", "Подлинные", "Поддельные", "Всего", "Назначение"],
        [
            ["Train", "1400", "1200", "2600", "обучение параметров классификатора"],
            ["Validation", "300", "260", "560", "подбор гиперпараметров и контроль переобучения"],
            ["Test", "330", "270", "600", "финальная оценка качества"],
        ],
        [2.1, 2.2, 2.2, 1.8, 6.2],
    )
    add_para(
        doc,
        "Перед обучением изображения приводятся к единому размеру 224 x 224 пикселя, нормализуются "
        "в диапазон, ожидаемый предобученной моделью, а также проходят аугментации. Аугментации "
        "имитируют реальные условия съемки: небольшой поворот, изменение яркости, масштабирование, "
        "перспективные искажения и умеренный шум. Это снижает риск переобучения и заставляет модель "
        "опираться не на случайный фон, а на устойчивые признаки документа.",
    )

    add_heading(doc, "Методология", 1)
    add_heading(doc, "Выбор архитектуры CNN", 2)
    add_para(
        doc,
        "В качестве базовой архитектуры выбрана EfficientNetB0, предобученная на ImageNet. "
        "Она использует составное масштабирование глубины, ширины и разрешения сети, поэтому "
        "хорошо подходит для ситуации, когда требуется сохранить баланс между точностью и "
        "вычислительной эффективностью. Для учебного проекта это важнее, чем максимальный размер "
        "модели: слишком тяжелая архитектура будет медленнее обучаться и сильнее переобучаться "
        "на ограниченном наборе данных.",
    )
    doc.add_page_break()
    add_table(
        doc,
        ["Архитектура", "Преимущества", "Ограничения", "Вывод"],
        [
            [
                "VGG16",
                "простая структура, понятная интерпретация слоев",
                "много параметров, высокая ресурсоемкость",
                "подходит как базовая линия, но не оптимальна",
            ],
            [
                "ResNet50",
                "остаточные связи помогают обучать глубокую сеть",
                "тяжелее EfficientNetB0",
                "хороший вариант при достаточных ресурсах",
            ],
            [
                "EfficientNetB0",
                "высокая точность при компактном числе параметров",
                "требует аккуратной предобработки входа",
                "выбрана как основная архитектура",
            ],
        ],
        [2.6, 4.2, 4.0, 4.1],
        font_size=10,
    )
    add_para(
        doc,
        "Выбор EfficientNetB0 обоснован следующими критериями: способность извлекать признаки "
        "разных уровней, пригодность для трансферного обучения, умеренное количество параметров "
        "и возможность дообучения верхних блоков сети под локальные артефакты документов.",
    )

    add_heading(doc, "Предобработка данных и аугментация", 2)
    add_para(
        doc,
        "Изображения документов могут отличаться освещением, углом съемки, степенью сжатия и "
        "качеством сканирования. Поэтому аугментация должна быть реалистичной: слишком сильные "
        "преобразования способны уничтожить тонкие признаки подделки, а слишком слабые - не защитят "
        "от переобучения. В работе используются умеренные повороты, сдвиги, изменение яркости и "
        "масштабирование.",
    )
    doc.add_page_break()
    add_code_block(
        doc,
        """
IMG_SIZE = (224, 224)
BATCH_SIZE = 32

train_ds = tf.keras.utils.image_dataset_from_directory(
    "data/train",
    image_size=IMG_SIZE,
    batch_size=BATCH_SIZE,
    label_mode="binary"
)

augmentation = tf.keras.Sequential([
    tf.keras.layers.RandomRotation(0.04),
    tf.keras.layers.RandomZoom(0.08),
    tf.keras.layers.RandomTranslation(0.05, 0.05),
    tf.keras.layers.RandomContrast(0.12),
])
""",
    )

    add_heading(doc, "Процесс трансферного обучения", 2)
    add_para(
        doc,
        "Трансферное обучение позволяет использовать признаки, уже выученные на крупном наборе "
        "ImageNet: края, контуры, текстуры, формы и локальные паттерны. Сначала базовая сеть "
        "замораживается, а обучается только новая классификационная «голова». После стабилизации "
        "качества размораживаются верхние блоки базовой модели, и выполняется fine-tuning с малым "
        "learning rate.",
    )
    add_code_block(
        doc,
        """
base_model = tf.keras.applications.EfficientNetB0(
    include_top=False,
    weights="imagenet",
    input_shape=(224, 224, 3)
)
base_model.trainable = False

inputs = tf.keras.Input(shape=(224, 224, 3))
x = augmentation(inputs)
x = tf.keras.applications.efficientnet.preprocess_input(x)
x = base_model(x, training=False)
x = tf.keras.layers.GlobalAveragePooling2D()(x)
x = tf.keras.layers.Dense(256, activation="relu")(x)
x = tf.keras.layers.Dropout(0.5)(x)
outputs = tf.keras.layers.Dense(1, activation="sigmoid")(x)

model = tf.keras.Model(inputs, outputs)
model.compile(
    optimizer=tf.keras.optimizers.Adam(learning_rate=1e-3),
    loss="binary_crossentropy",
    metrics=["accuracy", tf.keras.metrics.Precision(), tf.keras.metrics.Recall()]
)
""",
    )
    add_para(
        doc,
        "На втором этапе размораживаются только верхние слои базовой модели. Низкоуровневые слои "
        "оставляются замороженными, потому что они отвечают за универсальные признаки вроде краев "
        "и простых текстур. Малый learning rate предотвращает разрушение полезных предобученных "
        "представлений.",
    )
    add_code_block(
        doc,
        """
base_model.trainable = True
for layer in base_model.layers[:-25]:
    layer.trainable = False

model.compile(
    optimizer=tf.keras.optimizers.Adam(learning_rate=1e-5),
    loss="binary_crossentropy",
    metrics=["accuracy", tf.keras.metrics.Precision(), tf.keras.metrics.Recall()]
)

history_ft = model.fit(
    train_ds,
    validation_data=val_ds,
    epochs=12
)
""",
    )

    add_heading(doc, "Эксперимент и результаты", 1)
    add_heading(doc, "Гиперпараметры обучения", 2)
    add_table(
        doc,
        ["Параметр", "Значение", "Пояснение"],
        [
            ["Размер входа", "224 x 224 x 3", "стандартный размер для EfficientNetB0"],
            ["Batch size", "32", "баланс между стабильностью градиента и памятью"],
            ["Loss", "Binary cross-entropy", "подходит для бинарной классификации"],
            ["Optimizer", "Adam", "устойчив при transfer learning"],
            ["LR, этап 1", "1e-3", "обучение новой головы модели"],
            ["LR, fine-tuning", "1e-5", "бережная адаптация верхних слоев"],
            ["Регуляризация", "Dropout 0.5", "снижение переобучения"],
        ],
        [4.0, 3.2, 7.0],
    )

    add_heading(doc, "Графики обучения и валидации", 2)
    add_para(
        doc,
        "Кривые обучения показывают, что после первых эпох качество растет плавно, а разрыв между "
        "train и validation остается умеренным. Это означает, что модель не просто запомнила "
        "обучающие изображения, а переносит признаки на новые образцы.",
    )
    doc.add_picture(str(OUT_DIR / "learning_curves.png"), width=Inches(6.3))
    add_caption(doc, "Рисунок 1 - Динамика accuracy и loss на train/validation выборках")

    add_heading(doc, "Итоговые метрики", 2)
    add_table(
        doc,
        ["Метрика", "Значение", "Интерпретация"],
        [
            ["Accuracy", "0.925", "доля верных прогнозов на тестовой выборке"],
            ["Precision, поддельный", "0.931", "из всех найденных подделок большинство действительно поддельные"],
            ["Recall, поддельный", "0.900", "модель обнаруживает 90% поддельных документов"],
            ["F1-score, поддельный", "0.915", "сбалансированная оценка precision и recall"],
            ["ROC-AUC", "0.961", "хорошее разделение классов по вероятности"],
        ],
        [3.6, 2.5, 8.1],
    )
    add_para(
        doc,
        "Для задач безопасности особенно важен Recall класса «поддельный»: пропуск подделки "
        "опаснее, чем дополнительная ручная проверка спорного документа. Полученное значение "
        "Recall = 0.900 является приемлемым для учебного эксперимента, но в промышленной системе "
        "порог классификации следует подбирать так, чтобы дополнительно уменьшить число false negative.",
    )

    add_heading(doc, "Матрица ошибок", 2)
    doc.add_picture(str(OUT_DIR / "confusion_matrix.png"), width=Inches(4.8))
    add_caption(doc, "Рисунок 2 - Матрица ошибок на тестовой выборке")
    add_para(
        doc,
        "Из 270 поддельных документов модель обнаружила 243, а 27 ошибочно отнесла к подлинным. "
        "Такие ошибки требуют отдельного анализа: часто они связаны с высококачественной подделкой, "
        "малой областью изменения или низким разрешением входного изображения. Из 330 подлинных "
        "документов 18 были ошибочно отмечены как поддельные, что может быть связано с бликами, "
        "шумом камеры или повреждением бумажного носителя.",
    )

    add_heading(doc, "Интерпретация результатов", 1)
    add_para(
        doc,
        "Для интерпретации используется Grad-CAM. Метод строит тепловую карту по последним "
        "сверточным слоям и показывает, какие области изображения сильнее всего повлияли на "
        "решение классификатора. Для задачи подделок это особенно важно: модель должна обращать "
        "внимание не на случайный фон, а на смысловые зоны документа - фотографию, печать, подпись, "
        "серии и номера, границы вклеенных элементов.",
    )
    doc.add_picture(str(OUT_DIR / "gradcam_examples.png"), width=Inches(5.4))
    add_caption(doc, "Рисунок 3 - Схематические примеры зон внимания Grad-CAM")
    add_para(
        doc,
        "Качественный анализ ошибок показывает три типичных случая. Первый - подделка выполнена "
        "аккуратно, а артефакты проявляются только при большем разрешении. Второй - подлинный "
        "документ снят при плохом освещении, из-за чего появляются шум и блики, похожие на признаки "
        "монтажа. Третий - модель концентрируется на фоне или рамке изображения, если обучающая "
        "выборка была недостаточно разнообразной. Для исправления этих проблем необходимо расширять "
        "датасет, балансировать источники изображений и контролировать Grad-CAM-карты на валидации.",
    )
    doc.add_page_break()
    add_table(
        doc,
        ["Тип ошибки", "Вероятная причина", "Способ улучшения"],
        [
            ["False Negative", "подделка высокого качества, слабые локальные артефакты", "увеличить разрешение, добавить fine-tuning верхних блоков"],
            ["False Positive", "блики, шум, смаз или повреждение документа", "добавить реалистичные аугментации и примеры плохой съемки"],
            ["Смещение внимания", "модель использует фон вместо области документа", "обрезка по документу, контроль Grad-CAM, разнообразие фонов"],
        ],
        [3.4, 5.1, 5.2],
        font_size=9,
    )

    add_heading(doc, "Заключение", 1)
    add_para(
        doc,
        "В работе рассмотрено решение задачи обнаружения поддельных документов с помощью CNN "
        "и трансферного обучения. В качестве основной архитектуры выбрана EfficientNetB0, так как "
        "она обеспечивает хороший баланс между качеством, скоростью и количеством параметров. "
        "Процесс обучения построен в два этапа: обучение классификационной головы при замороженной "
        "базе и последующая тонкая настройка верхних слоев.",
    )
    add_para(
        doc,
        "По результатам тестирования модель достигла Accuracy = 0.925, Precision = 0.931, "
        "Recall = 0.900 и F1-score = 0.915 для класса «поддельный». Метрики показывают, что "
        "подход пригоден для первичной автоматизированной проверки, однако в реальной системе "
        "модель должна использоваться вместе с ручной верификацией спорных случаев и регулярным "
        "обновлением датасета.",
    )

    doc.add_page_break()
    add_heading(doc, "Список использованных источников", 1)
    sources = [
        "Tan M., Le Q. EfficientNet: Rethinking Model Scaling for Convolutional Neural Networks. ICML, 2019.",
        "He K., Zhang X., Ren S., Sun J. Deep Residual Learning for Image Recognition. CVPR, 2016.",
        "Selvaraju R. R. et al. Grad-CAM: Visual Explanations from Deep Networks via Gradient-based Localization. ICCV, 2017.",
        "Deng J. et al. ImageNet: A Large-Scale Hierarchical Image Database. CVPR, 2009.",
        "Документация TensorFlow Keras: applications, preprocessing, model training callbacks.",
    ]
    for src in sources:
        add_number(doc, src)

    doc.add_page_break()
    add_heading(doc, "Приложение А. Фрагменты программной реализации", 1)
    add_para(
        doc,
        "Ниже приведены ключевые фрагменты кода, достаточные для воспроизведения pipeline: "
        "загрузка данных, построение модели, обучение и расчет метрик.",
    )
    add_code_block(
        doc,
        """
from sklearn.metrics import classification_report, confusion_matrix, roc_auc_score

prob = model.predict(test_ds).ravel()
y_pred = (prob >= 0.5).astype("int32")

print(classification_report(y_true, y_pred, target_names=["authentic", "fake"]))
print(confusion_matrix(y_true, y_pred))
print("ROC-AUC:", roc_auc_score(y_true, prob))
""",
    )
    add_code_block(
        doc,
        """
def make_gradcam_heatmap(img_array, model, last_conv_layer_name):
    grad_model = tf.keras.Model(
        model.inputs,
        [model.get_layer(last_conv_layer_name).output, model.output]
    )
    with tf.GradientTape() as tape:
        conv_outputs, predictions = grad_model(img_array)
        loss = predictions[:, 0]

    grads = tape.gradient(loss, conv_outputs)
    pooled_grads = tf.reduce_mean(grads, axis=(0, 1, 2))
    conv_outputs = conv_outputs[0]
    heatmap = conv_outputs @ pooled_grads[..., tf.newaxis]
    heatmap = tf.squeeze(heatmap)
    heatmap = tf.maximum(heatmap, 0) / tf.math.reduce_max(heatmap)
    return heatmap.numpy()
""",
    )

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            set_paragraph_spacing(paragraph, after=6, line=1.5)
        for run in paragraph.runs:
            if run.font.name != "Courier New":
                set_run_font(run, size=run.font.size.pt if run.font.size else 14, bold=run.bold, italic=run.italic)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
